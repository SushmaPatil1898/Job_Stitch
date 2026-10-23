"""
parsing.py
----------
Turns whatever the user gives us (an uploaded PDF/DOCX/TXT file, or text
they typed/pasted into a box) into a single plain-text string we can run
analysis on.

Kept deliberately simple and dependency-light: pdfplumber for PDFs and
python-docx for Word files, both of which are pure "extract the text"
jobs with no formatting decisions to make.
"""

from __future__ import annotations

import io
from dataclasses import dataclass


class UnsupportedFileTypeError(Exception):
    """Raised when a file extension we don't know how to parse is uploaded."""


@dataclass
class ParsedDocument:
    text: str
    source_name: str
    char_count: int
    word_count: int


def _extract_pdf_text(file_bytes: bytes) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx_text(file_bytes: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]

    # Tables often hold skills/experience in resumes built from templates.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def parse_uploaded_file(uploaded_file) -> ParsedDocument:
    """
    uploaded_file: a Streamlit UploadedFile (has .name and .getvalue()),
    or any object exposing the same interface.
    """
    name = uploaded_file.name
    file_bytes = uploaded_file.getvalue()
    extension = name.lower().rsplit(".", 1)[-1] if "." in name else ""

    if extension == "pdf":
        text = _extract_pdf_text(file_bytes)
    elif extension in ("docx",):
        text = _extract_docx_text(file_bytes)
    elif extension in ("txt", "md"):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileTypeError(
            f"'.{extension}' files aren't supported. Please upload a PDF, "
            "DOCX, or TXT file, or paste the text directly instead."
        )

    text = text.strip()
    return ParsedDocument(
        text=text,
        source_name=name,
        char_count=len(text),
        word_count=len(text.split()),
    )


def parse_pasted_text(text: str, label: str = "pasted text") -> ParsedDocument:
    text = (text or "").strip()
    return ParsedDocument(
        text=text,
        source_name=label,
        char_count=len(text),
        word_count=len(text.split()),
    )


# Common resume section headers -- used as a stopping point for
# extract_header_lines() below, since the name/title/contact block always
# comes before the first of these.
_SECTION_KEYWORDS = {
    "summary", "objective", "profile", "about",
    "experience", "work experience", "employment", "employment history",
    "work history", "professional experience",
    "skills", "technical skills", "core competencies", "key skills",
    "education", "academic background",
    "certifications", "certificates", "licenses",
    "projects", "key projects",
}


def extract_header_lines(resume_text: str, max_lines: int = 4) -> list[str]:
    """
    Pulls out the first few lines of a resume -- name, job title, and
    contact details (phone/email/LinkedIn/location) -- which normally sit
    before the first real section. Used so that information always makes
    it into a rewritten/exported version of the resume, even if an AI
    rewrite step drops it.
    """
    header: list[str] = []
    for raw_line in resume_text.split("\n"):
        line = raw_line.strip()
        if not line:
            if header:
                break
            continue
        if line.lower().rstrip(":") in _SECTION_KEYWORDS:
            break
        header.append(line)
        if len(header) >= max_lines:
            break
    return header
